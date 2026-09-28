"""Multi-source document loaders - REFERENCE SOLUTION (Activity 1.2)."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from src.utils.config_loader import load_yaml_config, resolve_path
from src.utils.logger import setup_logger

logger = setup_logger(__name__)


@dataclass
class RawDocument:
    content: str
    source: str
    file_type: str


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(pages).strip()


def _load_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    paras = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    return "\n".join(paras).strip()


def _load_html(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _load_markdown(path: Path) -> str:
    return path.read_text(encoding="utf-8").strip()


_EXTENSION_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".html": _load_html,
    ".htm": _load_html,
    ".md": _load_markdown,
}


def load_document(path: str | Path, kb_root: str | Path = "data/knowledge_base") -> RawDocument:
    path = Path(path)
    kb_root = Path(kb_root)
    ext = path.suffix.lower()
    loader = _EXTENSION_LOADERS.get(ext)
    if loader is None:
        raise ValueError(f"Unsupported file type: {ext} ({path.name})")
    try:
        rel = str(path.relative_to(kb_root)).replace("\\", "/")
    except ValueError:
        rel = path.name
    return RawDocument(content=loader(path), source=rel, file_type=ext.lstrip("."))


def load_corpus(kb_root: str | Path = "data/knowledge_base",
                base_path: Optional[str | Path] = None) -> List[RawDocument]:
    kb_root = resolve_path(kb_root, base_path)
    supported = set(load_yaml_config("config/ingestion_config.yaml",
                                     base_path=base_path)["ingestion"]["supported_extensions"])
    docs: List[RawDocument] = []
    for file in sorted(kb_root.rglob("*")):
        if not file.is_file() or file.name == "manifest.csv":
            continue
        if file.suffix.lower() not in supported:
            logger.warning(f"Skipping unsupported file: {file}")
            continue
        docs.append(load_document(file, kb_root))
    logger.info(f"Loaded {len(docs)} documents from {kb_root}")
    return docs
